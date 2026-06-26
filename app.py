import io
import json
import os
import psycopg2         # replace with sqlite3 if needed
import psycopg2.extras
from psycopg2 import pool
import subprocess
import threading
import time
import uuid
import openpyxl
from collections import defaultdict
from pathlib import Path
from datetime import datetime, timezone
import re
import sys
from flask import session, redirect, url_for
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from contextlib import contextmanager

from dotenv import load_dotenv
load_dotenv()

connection_pool = pool.SimpleConnectionPool(
    minconn=1,
    maxconn=10,
    dsn=os.getenv("DATABASE_URL")
)


def _ensure_company_visibility_table():
    conn = connection_pool.getconn()
    conn.cursor_factory = psycopg2.extras.RealDictCursor
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS company_visibility (
            id SERIAL PRIMARY KEY,
            company_id INTEGER UNIQUE NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
            hide_company BOOLEAN DEFAULT FALSE,
            hide_employees BOOLEAN DEFAULT FALSE,
            hide_contact_info BOOLEAN DEFAULT FALSE
        )
    """)
    conn.commit()
    connection_pool.putconn(conn)

_ensure_company_visibility_table()

from db_addition import save_categories, save_flag, resolve_flag, get_flagged
from flask import Flask, Response, jsonify, render_template, request, stream_with_context

import requests

import mobile.script_mobile as mobile_ocr

# we can remove the taiga stuff. keeping it just in case we want to rollback
TAIGA_USERNAME = os.getenv("TAIGA_USERNAME")
TAIGA_PASSWORD = os.getenv("TAIGA_PASSWORD")
TAIGA_API_URL = "https://api.taiga.io/api/v1"

ZOHO_CLIENT_ID = os.getenv("ZOHO_CLIENT_ID")
ZOHO_CLIENT_SECRET = os.getenv("ZOHO_CLIENT_SECRET")
ZOHO_REFRESH_TOKEN = os.getenv("ZOHO_REFRESH_TOKEN")
ZOHO_BASE_URL = os.getenv("ZOHO_BASE_URL", "https://www.zohoapis.in/crm/v2")
ZOHO_ACCOUNTS_URL = os.getenv("ZOHO_ACCOUNTS_URL", "https://accounts.zoho.in")

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-secret-change-in-production")
app.config['SESSION_COOKIE_SECURE'] = True
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

# DB_PATH = os.environ.get("DB_PATH", str(Path(__file__).parent / "contacts.db"))

# ── live process registry ─────────────────────────────────────────────────────
_processes: dict[str, dict] = {}  # job_id → {proc, output, done}
_process_lock = threading.Lock()


# ─────────────────────────────────────────────────────────────────────────────
#  DB helpers
# ─────────────────────────────────────────────────────────────────────────────

def get_db():
    conn = connection_pool.getconn()
    conn.cursor_factory = psycopg2.extras.RealDictCursor
    return conn

def close_db(conn):
    connection_pool.putconn(conn)


@contextmanager
def db_ctx():
    """Yield (conn, cursor), roll back on exception, always return conn to pool."""
    conn = get_db()
    cursor = conn.cursor()
    try:
        yield conn, cursor
    except Exception:
        conn.rollback()
        raise
    finally:
        close_db(conn)


def safe_json(v):
    if not v or v == "null":
        return []
    try:
        parsed = json.loads(v)
        if parsed is None:
            return []
        if isinstance(parsed, list):
            return [i for i in parsed if i is not None]
        return [parsed]
    except Exception:
        return [v]


def get_companies_filtered(categories=None, industries=None, types=None, search=None, sort_by="name", sort_dir="asc"):
    sql = """
        SELECT
            c.id AS id,
            c.name,
            c.addresses,
            c.phones,
            c.emails,
            c.websites,
            c.company_info,
            c.source_folder,
            e.primary_industry,
            e.sub_industry,
            e.company_type,
            e.products,
            e.markets,
            e.domain,
            STRING_AGG(cc.category, '||') AS categories,
            c.notes
        FROM companies c
        LEFT JOIN enrichment e ON e.contact_id = c.id
        LEFT JOIN company_categories cc ON cc.company_id = c.id
        WHERE 1=1
    """
    params = []

    if search:
        sql += """
            AND (
                c.name ILIKE %s
                OR e.domain ILIKE %s
                OR c.company_info ILIKE %s
                OR EXISTS (
                    SELECT 1 FROM people p
                    WHERE p.company_id = c.id
                    AND (p.name ILIKE %s OR p.emails ILIKE %s)
                )
            )
        """
        params += [f"%{search}%"] * 5

    if industries:
        placeholders = ",".join(["%s"] * len(industries))
        sql += f" AND e.primary_industry IN ({placeholders})"
        params += industries

    if types:
        type_conditions = " OR ".join(["e.company_type LIKE %s" for _ in types])
        sql += f" AND ({type_conditions})"
        params += [f"%{t}%" for t in types]

    sql += " GROUP BY c.id, c.name, c.addresses, c.phones, c.websites, c.company_info, c.source_folder, c.notes, e.primary_industry, e.sub_industry, e.company_type, e.products, e.markets, e.domain"

    if categories:
        placeholders = ",".join(["%s"] * len(categories))
        sql = f"""
            SELECT sub.id, sub.name, sub.addresses, sub.phones, sub.emails, sub.websites,
                   sub.company_info, sub.source_folder, sub.primary_industry,
                   sub.sub_industry, sub.company_type, sub.products, sub.markets,
                   sub.domain, sub.categories, sub.notes
            FROM ({sql}) sub
            WHERE EXISTS (
                SELECT 1 FROM company_categories cc2
                WHERE cc2.company_id = sub.id
                AND cc2.category IN ({placeholders})
            )
        """
        params += categories

    valid_sorts = {"name", "primary_industry", "domain", "id"}
    if sort_by not in valid_sorts:
        sort_by = "name"
    sort_dir = "DESC" if sort_dir == "desc" else "ASC"
    sql += f" ORDER BY {sort_by} {sort_dir}"

    with db_ctx() as (_, cur):
        cur.execute(sql, params)
        rows = cur.fetchall()
    return rows


def get_all_categories():
    try:
        with db_ctx() as (_, cur):
            cur.execute("SELECT DISTINCT category FROM company_categories ORDER BY category")
            return [r["category"] for r in cur.fetchall()]
    except Exception:
        return []


def get_all_industries():
    try:
        with db_ctx() as (_, cur):
            cur.execute("SELECT DISTINCT primary_industry FROM enrichment WHERE primary_industry IS NOT NULL ORDER BY primary_industry")
            return [r["primary_industry"] for r in cur.fetchall()]
    except Exception:
        return []


def get_all_types():
    try:
        with db_ctx() as (_, cur):
            cur.execute("SELECT company_type FROM enrichment WHERE company_type IS NOT NULL")
            rows = cur.fetchall()
    except Exception:
        return []

    seen = set()
    for row in rows:
        try:
            types = json.loads(row["company_type"])
            if isinstance(types, list):
                for t in types:
                    if t and isinstance(t, str):
                        seen.add(t.strip())
            elif isinstance(types, str) and types:
                seen.add(types.strip())
        except (json.JSONDecodeError, TypeError):
            pass
    return sorted(seen)


def get_stats():
    try:
        with db_ctx() as (_, cur):
            cur.execute("SELECT COUNT(*) as n FROM companies")
            companies = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) as n FROM people")
            people = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) as n FROM enrichment WHERE domain IS NOT NULL")
            enriched = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(DISTINCT category) as n FROM company_categories")
            categories = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) as n FROM flagged WHERE status = 'active'")
            flagged = cur.fetchone()["n"]
        return {"companies": companies, "people": people, "enriched": enriched,
                "categories": categories, "flagged": flagged}
    except Exception:
        return {"companies": 0, "people": 0, "enriched": 0, "categories": 0, "flagged": 0}


def get_people_for_company(company_id):
    with db_ctx() as (_, cur):
        cur.execute("SELECT * FROM people WHERE company_id = %s", (company_id,))
        return cur.fetchall()

# ─────────────────────────────────────────────────────────────────────────────
# Taiga helpers
# ─────────────────────────────────────────────────────────────────────────────

# def get_taiga_token():
#     res = requests.post(f"{TAIGA_API_URL}/auth", json={
#         "type": "normal",
#         "username": TAIGA_USERNAME,
#         "password": TAIGA_PASSWORD,
#     })
#     data = res.json()
#     return data.get("auth_token")


# # def get_taiga_project_id(token):
# #     res = requests.get(
# #         f"{TAIGA_API_URL}/projects/by_slug%sslug={TAIGA_PROJECT_SLUG}",
# #         headers={"Authorization": f"Bearer {token}"}
# #     )
# #     return res.json().get("id")


# def taiga_create_issue(title, description, project_id, due_date=None):
#     token = get_taiga_token()
#     if not token:
#         return None, "Failed to authenticate with Taiga", None

#     payload = {
#         "project": project_id,
#         "subject": title,
#         "description": description or "",
#     }
#     if due_date:
#         payload["due_date"] = due_date

#     res = requests.post(
#         f"{TAIGA_API_URL}/userstories",
#         json=payload,
#         headers={"Authorization": f"Bearer {token}"}
#     )
#     data = res.json()
#     if res.status_code == 201:
#         return data.get("ref"), None, data.get("id")
#     else:
#         return None, data, None

# # Slug helpers

# def make_project_slug(company_name):
#     slug = company_name.lower()
#     slug = re.sub(r'[^a-z0-9]+', '-', slug)
#     slug = slug.strip('-')
#     return f"yl-{slug}"


# def get_or_create_taiga_project(company_id, company_name, token):
    # conn = get_db()
    # row = conn.execute(
    #     "SELECT taiga_project_id, taiga_project_slug FROM companies WHERE id = %s",
    #     (company_id,)
    # ).fetchone()
    # close_db(conn)

    # if row and row["taiga_project_id"] and row["taiga_project_slug"]:
    #     return row["taiga_project_id"], row["taiga_project_slug"], None

    # # create new project
    # res = requests.post(
    #     f"{TAIGA_API_URL}/projects",
    #     json={
    #         "name": company_name,
    #         "description": f"YantraLive project for {company_name}",
    #         "is_private": True,
    #     },
    #     headers={"Authorization": f"Bearer {token}"}
    # )

    # if res.status_code != 201:
    #     return None, None, f"Failed to create Taiga project: {res.text}"

    # data = res.json()
    # project_id = data.get("id")
    # project_slug = data.get("slug")

    # conn = get_db()
    # conn.execute(
    #     "UPDATE companies SET taiga_project_id = %s, taiga_project_slug = %s WHERE id = %s",
    #     (project_id, project_slug, company_id)
    # )
    # conn.commit()
    # close_db(conn)

    # return project_id, project_slug, None

# ─────────────────────────────────────────────────────────────────────────────
# Zoho CRM helpers
# ─────────────────────────────────────────────────────────────────────────────
def get_zoho_access_token():
    res = requests.post(f"{ZOHO_ACCOUNTS_URL}/oauth/v2/token", data={
        "grant_type": "refresh_token",
        "client_id": ZOHO_CLIENT_ID,
        "client_secret": ZOHO_CLIENT_SECRET,
        "refresh_token": ZOHO_REFRESH_TOKEN,
    })
    data = res.json()
    return data.get("access_token")


def zoho_search_account(name, token):
    """Search for an existing Account by name, return its ID or None."""
    res = requests.get(
        f"{ZOHO_BASE_URL}/Accounts/search",
        params={"criteria": f"Account_Name:equals:{name}"},
        headers={"Authorization": f"Zoho-oauthtoken {token}"}
    )
    if res.status_code == 200:
        data = res.json().get("data", [])
        if data:
            return data[0]["id"]
    return None


def zoho_create_task(title, description, due_date, account_id, token):
    """Create a Task in Zoho CRM, optionally linked to an Account."""
    payload = {
        "data": [{
            "Subject": title,
            "Description": description or "",
            "Status": "Not Started",
            "Due_Date": due_date or None,
        }]
    }
    if account_id:
        payload["data"][0]["What_Id"] = account_id
        payload["data"][0]["$se_module"] = "Accounts"

    res = requests.post(
        f"{ZOHO_BASE_URL}/Tasks",
        json=payload,
        headers={"Authorization": f"Zoho-oauthtoken {token}"}
    )
    return res.status_code in (200, 201), res.json()

# ─────────────────────────────────────────────────────────────────────────────
# User-agent check for mobile
# ─────────────────────────────────────────────────────────────────────────────

def is_mobile():
    ua = request.headers.get('User-Agent', '').lower()
    mobile_keywords = ['mobile', 'android', 'iphone', 'ipad', 'ipod', 'blackberry', 'windows phone']
    return any(kw in ua for kw in mobile_keywords)


# ────────────────────────────────────────────────────────────────────
# Auth routes
# ────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login_page"))
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login_page"))
        if session.get("role") not in ("admin", "superadmin"):
            return jsonify({"error": "admin only"}), 403
        return f(*args, **kwargs)
    return decorated

@app.route("/login")
def login_page():
    if "user_id" in session:
        return redirect(url_for("index"))
    return render_template("login.html")

@app.route("/api/auth/login", methods=["POST"])
def api_login():
    data = request.json or {}
    username = data.get("username", "").strip()
    password = data.get("password", "")

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE username = %s", (username,))
    user = cursor.fetchone()
    close_db(conn)

    if not user or not check_password_hash(user["password_hash"], password):
        return jsonify({"error": "Invalid username or password"}), 401

    session["user_id"] = user["id"]
    session["username"] = user["username"]
    session["role"] = user["role"]
    return jsonify({"ok": True, "role": user["role"]})

@app.route("/api/auth/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"ok": True})

@app.route("/api/auth/me")
def api_me():
    if "user_id" not in session:
        return jsonify({"logged_in": False})
    return jsonify({
        "logged_in": True,
        "user_id": session["user_id"],
        "username": session["username"],
        "role": session["role"]
    })

@app.route("/admin")
@admin_required
def admin_page():
    return render_template("admin.html")

@app.route("/api/admin/users", methods=["GET"])
@admin_required
def api_get_users():
    with db_ctx() as (_, cur):
        cur.execute("SELECT id, username, role, created_at FROM users ORDER BY created_at DESC")
        rows = cur.fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/admin/users", methods=["POST"])
@admin_required
def api_create_user():
    data = request.json or {}
    username = data.get("username", "").strip()
    password = data.get("password", "")
    role = data.get("role", "user")

    if not username or not password:
        return jsonify({"error": "username and password required"}), 400

    if role == "superadmin":
        return jsonify({"error": "superadmin accounts cannot be created from the dashboard"}), 403

    try:
        with db_ctx() as (conn, cur):
            cur.execute("""
                INSERT INTO users (username, password_hash, role, created_at)
                VALUES (%s, %s, %s, %s)
            """, (username, generate_password_hash(password), role, datetime.now().isoformat()))
            conn.commit()
    except Exception:
        return jsonify({"error": "username already exists"}), 409
    return jsonify({"ok": True})

@app.route("/api/admin/users/<int:user_id>", methods=["DELETE"])
@admin_required
def api_delete_user(user_id):
    if user_id == session["user_id"]:
        return jsonify({"error": "You cannot delete your own account"}), 400

    with db_ctx() as (conn, cur):
        cur.execute("SELECT role FROM users WHERE id = %s", (user_id,))
        user = cur.fetchone()

        if user and user["role"] == "superadmin":
            return jsonify({"error": "Superadmin accounts cannot be deleted from the dashboard"}), 403

        cur.execute("SELECT COUNT(*) as count FROM users WHERE role = 'admin'")
        admin_count = cur.fetchone()["count"]

        if user and user["role"] == "admin" and admin_count <= 1:
            return jsonify({"error": "Cannot delete the last admin account"}), 400

        cur.execute("DELETE FROM users WHERE id = %s", (user_id,))
        conn.commit()
    return jsonify({"ok": True})

@app.route("/api/admin/export", methods=["GET"])
@admin_required
def api_export_db():
    TABLES = [
        ("companies",          "SELECT * FROM companies ORDER BY id"),
        ("people",             "SELECT * FROM people ORDER BY id"),
        ("enrichment",         "SELECT * FROM enrichment ORDER BY id"),
        ("company_categories", "SELECT * FROM company_categories ORDER BY id"),
        ("tasks",              "SELECT * FROM tasks ORDER BY id"),
    ]

    wb = openpyxl.Workbook()
    if wb.active is not None:
        wb.remove(wb.active)

    with db_ctx() as (_, cur):
        for table_name, query in TABLES:
            cur.execute(query)
            rows = cur.fetchall()
            ws = wb.create_sheet(title=table_name)
            if rows:
                ws.append(list(rows[0].keys()))
                for row in rows:
                    ws.append(list(row.values()))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"yantralive_export_{timestamp}.xlsx"
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ─────────────────────────────────────────────────────────────────────────────
#  Routes
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/")
@login_required
def index():
    stats = get_stats()
    categories = get_all_categories()
    industries = get_all_industries()
    types = get_all_types()
    if is_mobile():
        return render_template("mobile.html", stats=stats, categories=categories, industries=industries, types=types)
    return render_template("index.html", stats=stats, categories=categories, industries=industries, types=types)


@app.route("/api/companies")
@login_required
def api_companies():
    cats = request.args.getlist("cat")
    industries = request.args.getlist("industry")
    types = request.args.getlist("type")
    search = request.args.get("q", "").strip() or None
    sort_by = request.args.get("sort", "name")
    sort_dir = request.args.get("dir", "asc")

    rows = get_companies_filtered(
        categories=cats or None,
        industries=industries or None,
        types=types or None,
        search=search,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )

    is_admin = session.get("role") in ("admin", "superadmin")

    # Fetch visibility settings for all returned companies
    company_ids = [r["id"] for r in rows]
    visibility_map = {}
    if company_ids:
        with db_ctx() as (_, cur):
            cur.execute("""
                SELECT company_id, hide_company, hide_employees, hide_contact_info
                FROM company_visibility
                WHERE company_id = ANY(%s)
            """, (company_ids,))
            for v in cur.fetchall():
                visibility_map[v["company_id"]] = dict(v)

    result = []
    for r in rows:
        cats_list = list(set(r["categories"].split("||"))) if r["categories"] else []
        vis = visibility_map.get(r["id"], {})
        hide_company = vis.get("hide_company", False)
        hide_employees = vis.get("hide_employees", False)
        hide_contact_info = vis.get("hide_contact_info", False)

        # Non-admins skip fully hidden companies
        if not is_admin and hide_company:
            continue

        phones = safe_json(r["phones"])
        emails = safe_json(r.get("emails"))

        company = {
            "id": r["id"],
            "name": r["name"],
            "domain": r["domain"],
            "primary_industry": r["primary_industry"],
            "sub_industry": r["sub_industry"],
            "company_type": safe_json(r["company_type"]),
            "products": safe_json(r["products"]),
            "markets": safe_json(r["markets"]),
            "categories": sorted(cats_list),
            "websites": safe_json(r["websites"]),
            "phones": [] if (not is_admin and hide_contact_info) else phones,
            "emails": [] if (not is_admin and hide_contact_info) else emails,
            "addresses": safe_json(r["addresses"]),
            "notes": r["notes"] or "",
        }

        if is_admin:
            company["visibility"] = {
                "hide_company": hide_company,
                "hide_employees": hide_employees,
                "hide_contact_info": hide_contact_info,
            }
        else:
            company["employees_hidden"] = hide_employees
            company["contact_info_hidden"] = hide_contact_info

        result.append(company)
    return jsonify(result)


@app.route("/api/company/<int:company_id>/people")
@login_required
def api_people(company_id):
    is_admin = session.get("role") in ("admin", "superadmin")

    hide_employees = False
    hide_contact_info = False
    if not is_admin:
        with db_ctx() as (_, cur):
            cur.execute("""
                SELECT hide_employees, hide_contact_info
                FROM company_visibility WHERE company_id = %s
            """, (company_id,))
            vis = cur.fetchone()
        if vis:
            hide_employees = vis["hide_employees"]
            hide_contact_info = vis["hide_contact_info"]

    if hide_employees:
        return jsonify({"hidden": True, "people": []})

    people = get_people_for_company(company_id)
    result = [{
        "id": p["id"],
        "company_id": company_id, 
        "name": p["name"],
        "role": p["role"],
        "phones": [] if hide_contact_info else safe_json(p["phones"]),
        "emails": [] if hide_contact_info else safe_json(p["emails"]),
        "notes": p["notes"] or ""
    } for p in people]
    return jsonify({"hidden": False, "people": result})


@app.route("/api/stats")
@login_required
def api_stats():
    return jsonify(get_stats())


# @app.route("/api/company/<int:company_id>/notes", methods=["POST"])
# @login_required
# def save_notes(company_id):
#     notes = (request.json or {}).get("notes", "")
#     with db_ctx() as (conn, cur):
#         cur.execute("UPDATE companies SET notes = %s WHERE id = %s", (notes, company_id))
#         conn.commit()
#     return jsonify({"ok": True})

@app.route("/api/company/<int:company_id>/domain", methods=["POST"])
@admin_required
def save_domain(company_id):
    domain = (request.json or {}).get("domain", "").strip()
    if not domain:
        return jsonify({"error": "no domain provided"}), 400
    domain = domain.replace("https://", "").replace("http://", "").rstrip("/")
    with db_ctx() as (conn, cur):
        cur.execute("UPDATE companies SET websites = %s WHERE id = %s",
                    (json.dumps([domain]), company_id))
        conn.commit()
    resolve_flag(company_id, "no_domain", "resolved")
    return jsonify({"ok": True})

@app.route("/api/transcript", methods=["POST"])
@login_required
def process_transcript():
    data = request.json or {}
    transcript = data.get("transcript", "").strip()
    if not transcript:
        return jsonify({"error": "no transcript provided"}), 400

    import openai
    client = openai.OpenAI()
    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{
            "role": "user",
            "content": f"""You are extracting action items from a meeting transcript.
Return a JSON object with two keys:
- "company": the company name mentioned most in the context of action items (or null if unclear)
- "tasks": list of objects each with: title, description, assignee, due_date (YYYY-MM-DD or null)

Return only JSON, no explanation.

Transcript:
{transcript[:8000]}"""
        }]
    )
    raw = response.choices[0].message.content
    clean = raw.replace("```json", "").replace("```", "").strip()
    try:
        result = json.loads(clean)
    except json.JSONDecodeError:
        return jsonify({"error": "failed to parse GPT response"}), 500

    # fuzzy match company name to DB
    suggested_company = None
    suggested_company_id = None
    gpt_company = result.get("company")
    if gpt_company:
        with db_ctx() as (_, cur):
            cur.execute("SELECT id, name FROM companies")
            rows = cur.fetchall()
        gpt_lower = gpt_company.lower()
        best_match = None
        best_score = 0
        for row in rows:
            name = row["name"] or ""
            # simple overlap score
            score = sum(1 for word in gpt_lower.split() if word in name.lower())
            if score > best_score:
                best_score = score
                best_match = row
        if best_match and best_score > 0:
            suggested_company = best_match["name"]
            suggested_company_id = best_match["id"]

    return jsonify({
        "tasks": result.get("tasks", []),
        "suggested_company": suggested_company,
        "suggested_company_id": suggested_company_id,
    })


@app.route("/api/tasks", methods=["POST"])
@login_required
def save_tasks():
    data = request.json or {}
    company_id = data.get("company_id")
    tasks = data.get("tasks", [])
    transcript = data.get("transcript", "")

    with db_ctx() as (conn, cur):
        for task in tasks:
            cur.execute("""
                INSERT INTO tasks (company_id, title, description, assignee, due_date, status, source_transcript, created_at)
                VALUES (%s, %s, %s, %s, %s, 'To Do', %s, %s)
            """, (
                company_id,
                task.get("title"),
                task.get("description"),
                task.get("assignee"),
                task.get("due_date"),
                transcript,
                datetime.now().isoformat()
            ))
        conn.commit()
    return jsonify({"ok": True, "saved": len(tasks)})

@app.route("/api/tasks/<int:task_id>", methods=["DELETE"])
@login_required
def delete_task(task_id):
    with db_ctx() as (conn, cur):
        cur.execute("DELETE FROM tasks WHERE id = %s", (task_id,))
        conn.commit()
    return jsonify({"ok": True})


@app.route("/api/tasks", methods=["GET"])
@login_required
def get_tasks():
    company_id = request.args.get("company_id")
    status = request.args.get("status")
    sql = """
        SELECT t.*, c.name as company_name
        FROM tasks t
        LEFT JOIN companies c ON c.id = t.company_id
        WHERE 1=1
    """
    params = []
    if company_id:
        sql += " AND t.company_id = %s"
        params.append(company_id)
    if status:
        sql += " AND t.status = %s"
        params.append(status)
    sql += " ORDER BY t.created_at DESC"
    with db_ctx() as (_, cur):
        cur.execute(sql, params)
        rows = cur.fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/tasks/<int:task_id>", methods=["PATCH"])
@login_required
def update_task(task_id):
    data = request.json or {}
    allowed = {"title", "description", "assignee", "due_date", "status"}
    updates = {k: v for k, v in data.items() if k in allowed}
    if not updates:
        return jsonify({"error": "nothing to update"}), 400
    set_clause = ", ".join(f"{k} = %s" for k in updates)
    values = list(updates.values()) + [task_id]
    with db_ctx() as (conn, cur):
        cur.execute(f"UPDATE tasks SET {set_clause} WHERE id = %s", values)
        conn.commit()
    return jsonify({"ok": True})


@app.route("/api/companies/list")
@login_required
def api_companies_list():
    with db_ctx() as (_, cur):
        cur.execute("SELECT id, name FROM companies ORDER BY name")
        rows = cur.fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/enrich/single", methods=["POST"])
@admin_required
def enrich_single():
    company_id = (request.json or {}).get("company_id")
    if not company_id:
        return jsonify({"error": "company_id required"}), 400

    job_id = str(uuid.uuid4())[:8]

    def run():
        cmd = [sys.executable, "enrich2.py", "--single", str(company_id)]
        _stream_process(job_id, cmd)

    t = threading.Thread(target=run, daemon=True)
    t.start()
    return jsonify({"job_id": job_id})

# ─────────────────────────────────────────────────────────────────────────────
# Taiga routes
# ─────────────────────────────────────────────────────────────────────────────

# some taiga thing

# @app.route("/api/tasks/<int:task_id>/push-taiga", methods=["POST"])
# def push_to_taiga(task_id):
#     conn = get_db()
#     task = conn.execute("""
#         SELECT t.*, c.name as company_name
#         FROM tasks t
#         LEFT JOIN companies c ON c.id = t.company_id
#         WHERE t.id = %s
#     """, (task_id,)).fetchone()
#     close_db(conn)

#     if not task:
#         return jsonify({"error": "task not found"}), 404

#     token = get_taiga_token()
#     if not token:
#         return jsonify({"error": "Failed to authenticate with Taiga"}), 500

#     project_id, project_slug, error = get_or_create_taiga_project(
#         task["company_id"], task["company_name"], token
#     )
#     if error:
#         return jsonify({"error": error}), 500

#     ref, error, taiga_id = taiga_create_issue(
#         title=task["title"],
#         description=task["description"],
#         project_id=project_id,
#         due_date=task["due_date"],
#     )
#     if error:
#         return jsonify({"error": error}), 500

#     url = f"https://tree.taiga.io/project/{project_slug}/us/{ref}"
#     return jsonify({"ok": True, "url": url, "ref": ref})

# @app.route("/api/tasks/<int:task_id>/push-taiga", methods=["POST"])
# def push_to_taiga(task_id):
#     conn = get_db()
#     task = conn.execute("""
#         SELECT t.*, c.name as company_name
#         FROM tasks t
#         LEFT JOIN companies c ON c.id = t.company_id
#         WHERE t.id = %s
#     """, (task_id,)).fetchone()
#     close_db(conn)

#     if not task:
#         return jsonify({"error": "task not found"}), 404

#     token = get_taiga_token()
#     if not token:
#         return jsonify({"error": "Failed to authenticate with Taiga"}), 500

#     project_id, project_slug, error = get_or_create_taiga_project(
#         task["company_id"], task["company_name"], token
#     )
#     if error:
#         return jsonify({"error": error}), 500

#     ref, error, taiga_id = taiga_create_issue(
#         title=task["title"],
#         description=task["description"],
#         project_id=project_id,
#         due_date=task["due_date"],
#     )
#     if error:
#         return jsonify({"error": error}), 500

#     url = f"https://tree.taiga.io/project/{project_slug}/us/{ref}"
#     return jsonify({"ok": True, "url": url, "ref": ref})

# ─────────────────────────────────────────────────────────────────────────────
# zoho routes
# ─────────────────────────────────────────────────────────────────────────────


@app.route("/api/tasks/<int:task_id>/push-zoho", methods=["POST"])
@login_required
def push_to_zoho(task_id):
    with db_ctx() as (_, cur):
        cur.execute("""
            SELECT t.*, c.name as company_name
            FROM tasks t
            LEFT JOIN companies c ON c.id = t.company_id
            WHERE t.id = %s
        """, (task_id,))
        task = cur.fetchone()

    if not task:
        return jsonify({"error": "task not found"}), 404

    token = get_zoho_access_token()
    if not token:
        return jsonify({"error": "Failed to get Zoho access token"}), 500

    # search for matching account
    account_id = None
    if task["company_name"]:
        account_id = zoho_search_account(task["company_name"], token)

    success, result = zoho_create_task(
        title=task["title"],
        description=task["description"],
        due_date=task["due_date"],
        account_id=account_id,
        token=token,
    )

    if success:
        return jsonify({"ok": True, "account_matched": account_id is not None})
    else:
        return jsonify({"error": result}), 500
    
# ─────────────────────────────────────────────────────────────────────────────
#  Mobile routes
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/mobile")
@login_required
def mobile_index():
    return render_template("mobile.html")


@app.route("/mobile/scan", methods=["POST"])
@login_required
def mobile_scan():
    files = request.files.getlist("images")
    if not files or len(files) == 0:
        return jsonify({"error": "no images provided"}), 400
    if len(files) > 2:
        return jsonify({"error": "maximum 2 images (front and back)"}), 400

    images = []
    for f in files:
        mime_type = mobile_ocr.get_mime_type(f.filename)
        if not mime_type:
            return jsonify({"error": f"unsupported file type: {f.filename}"}), 400
        images.append((f.read(), mime_type))

    run_id = mobile_ocr.make_run_id()
    data = mobile_ocr.read_business_card_bytes(images, run_id)

    if not data:
        return jsonify({"error": "failed to extract data from image"}), 500

    return jsonify(data)


@app.route("/mobile/save", methods=["POST"])
@login_required
def mobile_save():
    data = request.json or {}
    if not data:
        return jsonify({"error": "no data provided"}), 400

    import db_addition as newdb
    newdb.init_db()
    company_id = newdb.save_company(data, "mobile")
    people_ids = newdb.save_people(company_id, data)

    return jsonify({
        "ok": True,
        "company_id": company_id,
        "people_ids": people_ids,
        "people_saved": len(data.get("people", []))
    })


@app.route("/mobile/enrich/<int:company_id>", methods=["POST"])
@login_required
def mobile_enrich(company_id):
    import threading
    def run():
        from enrich2 import scrape_website, enrich_company, save_enrichment, is_already_enriched
        import json
        with db_ctx() as (_, cur):
            cur.execute("SELECT websites FROM companies WHERE id = %s", (company_id,))
            row = cur.fetchone()
            cur.execute("SELECT MIN(emails) as emails FROM people WHERE company_id = %s", (company_id,))
            people_row = cur.fetchone()

        if is_already_enriched(company_id):
            return

        domain = None
        if row and row["websites"]:
            try:
                websites = json.loads(row["websites"])
                if isinstance(websites, list) and websites:
                    domain = websites[0].replace("https://","").replace("http://","").rstrip("/")
            except: pass

        if not domain and people_row and people_row["emails"]:
            from enrich2 import extract_domain
            domain = extract_domain(people_row["emails"])

        if not domain:
            return

        website_text = scrape_website(domain, "mobile")
        if not website_text:
            return

        enrichment = enrich_company(website_text, "mobile")
        save_enrichment(company_id, domain, enrichment)

    threading.Thread(target=run, daemon=True).start()
    return jsonify({"ok": True, "started": True})


@app.route("/mobile/enrich/<int:company_id>", methods=["GET"])
@login_required
def mobile_enrich_status(company_id):
    row = None
    try:
        with db_ctx() as (_, cur):
            cur.execute("SELECT * FROM enrichment WHERE contact_id = %s", (company_id,))
            row = cur.fetchone()
    except Exception:
        pass

    if not row:
        return jsonify({"status": "pending"})

    return jsonify({
        "status": "done",
        "primary_industry": row["primary_industry"],
        "sub_industry": row["sub_industry"],
        "company_type": safe_json(row["company_type"]),
        "products": safe_json(row["products"]),
        "markets": safe_json(row["markets"]),
        "domain": row["domain"]
    })

# ─────────────────────────────────────────────────────────────────────────────
#  Flagged routes
# ─────────────────────────────────────────────────────────────────────────────

CANONICAL_CATEGORIES = [
    "Earthmoving Equipment",
    "Concrete & Paving Equipment",
    "Lifting & Material Handling",
    "Crushing & Screening",
    "Drilling & Piling Equipment",
    "Mining Equipment",
    "Power Generation & Compressors",
    "Road & Infrastructure Construction",
    "Spare Parts & Components",
    "Tyres & Wheels",
    "Hydraulic Systems & Attachments",
    "Lubricants & Fluids",
    "Telematics & Machine Control",
    "Logistics & Transport",
    "Finance & Services",
    "Raw & Construction Materials",
    "Automobiles & Vehicles",
]


@app.route("/api/flagged")
@admin_required
def api_flagged():
    rows = get_flagged()
    result = []
    for r in rows:
        products = []
        if r["products"]:
            try:
                parsed = json.loads(r["products"])
                if isinstance(parsed, list):
                    products = [p for p in parsed if p]
            except Exception:
                pass
        result.append({
            "flag_id": r["flag_id"],
            "company_id": r["company_id"],
            "company_name": r["company_name"],
            "reason": r["reason"],
            "created_at": r["created_at"],
            "domain": r["domain"],
            "products": products,
        })
    return jsonify(result)


@app.route("/api/flagged/categories")
@admin_required
def api_flag_categories():
    return jsonify(CANONICAL_CATEGORIES)


@app.route("/api/flagged/<int:flag_id>/dismiss", methods=["POST"])
@admin_required
def api_dismiss_flag(flag_id):
    with db_ctx() as (_, cur):
        cur.execute("SELECT company_id, reason FROM flagged WHERE id = %s", (flag_id,))
        row = cur.fetchone()
    if not row:
        return jsonify({"error": "flag not found"}), 404
    resolve_flag(row["company_id"], row["reason"], "dismissed")
    return jsonify({"ok": True})


@app.route("/api/flagged/<int:flag_id>/assign", methods=["POST"])
@admin_required
def api_assign_flag(flag_id):
    data = request.json or {}
    categories = data.get("categories", [])
    if not categories:
        return jsonify({"error": "no categories provided"}), 400

    with db_ctx() as (_, cur):
        cur.execute("SELECT company_id, reason FROM flagged WHERE id = %s", (flag_id,))
        row = cur.fetchone()
    if not row:
        return jsonify({"error": "flag not found"}), 404

    save_categories(row["company_id"], categories)
    resolve_flag(row["company_id"], row["reason"], "resolved")
    return jsonify({"ok": True})


@app.route("/api/flagged/<int:flag_id>/gpt", methods=["POST"])
@admin_required
def api_gpt_suggest(flag_id):
    with db_ctx() as (_, cur):
        cur.execute("""
            SELECT f.company_id, f.reason, e.products
            FROM flagged f
            LEFT JOIN enrichment e ON e.contact_id = f.company_id
            WHERE f.id = %s
        """, (flag_id,))
        row = cur.fetchone()
    if not row:
        return jsonify({"error": "flag not found"}), 404

    products = []
    if row["products"]:
        try:
            parsed = json.loads(row["products"])
            if isinstance(parsed, list):
                products = [p for p in parsed if p]
        except Exception:
            pass

    if not products:
        return jsonify({"error": "no products to classify"}), 400

    try:
        import openai
        client = openai.OpenAI()
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{
                "role": "user",
                "content": f"""You are classifying a company into product categories for a B2B heavy equipment marketplace serving MENA and Africa.

The company's products/services are:
{json.dumps(products, indent=2)}

The available categories are:
{json.dumps(CANONICAL_CATEGORIES, indent=2)}

Return a JSON array of the category names that best match this company. Only include categories that are a genuine match. Return only JSON, no explanation."""
            }],
        )
        raw = response.choices[0].message.content
        clean = raw.replace("```json", "").replace("```", "").strip()
        suggested = json.loads(clean)
        if not isinstance(suggested, list):
            suggested = []
        # filter to only valid canonical categories
        suggested = [c for c in suggested if c in CANONICAL_CATEGORIES]
        return jsonify({"categories": suggested})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/flagged/gpt-bulk", methods=["POST"])
@admin_required
def api_gpt_bulk():
    rows = get_flagged()
    # only process no_category_match flags since others don't have products to classify
    to_process = [r for r in rows if r["reason"] == "no_category_match"]

    results = {"resolved": 0, "failed": 0, "skipped": 0}

    try:
        import openai
        client = openai.OpenAI()
    except Exception as e:
        return jsonify({"error": f"OpenAI init failed: {e}"}), 500

    for row in to_process:
        products = []
        if row["products"]:
            try:
                parsed = json.loads(row["products"])
                if isinstance(parsed, list):
                    products = [p for p in parsed if p]
            except Exception:
                pass

        if not products:
            results["skipped"] += 1
            continue

        try:
            response = client.chat.completions.create(
                model="gpt-4o",
                messages=[{
                    "role": "user",
                    "content": f"""You are classifying a company into product categories for a B2B heavy equipment marketplace serving MENA and Africa.

                    The company's products/services are:
                    {json.dumps(products, indent=2)}

                    The available categories are:
                    {json.dumps(CANONICAL_CATEGORIES, indent=2)}

                    Return a JSON array of the category names that best match this company. Only include categories that are a genuine match. Return only JSON, no explanation."""
                }],
            )
            raw = response.choices[0].message.content
            clean = raw.replace("```json", "").replace("```", "").strip()
            suggested = json.loads(clean)
            if not isinstance(suggested, list):
                suggested = []
            suggested = [c for c in suggested if c in CANONICAL_CATEGORIES]

            if suggested:
                save_categories(row["company_id"], suggested)
                resolve_flag(row["company_id"], "no_category_match", "resolved")
                results["resolved"] += 1
            else:
                results["skipped"] += 1
        except Exception:
            results["failed"] += 1

    return jsonify(results)

# ─────────────────────────────────────────────────────────────────────────────
#  Person notes
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/api/people/<int:person_id>/notes", methods=["GET"])
@login_required
def get_person_notes(person_id):
    role = session.get("role")
    user_id = session["user_id"]

    with db_ctx() as (_, cur):
        if role in ("admin", "superadmin"):
            cur.execute("""
                SELECT pn.id, pn.note, pn.created_at, pn.visibility, u.username
                FROM person_notes pn
                JOIN users u ON u.id = pn.user_id
                WHERE pn.person_id = %s
                ORDER BY pn.created_at DESC
            """, (person_id,))
            notes = [dict(r) for r in cur.fetchall()]
            hidden_count = 0
        else:
            cur.execute("""
                SELECT g.note_id FROM note_access_grants g
                JOIN note_access_requests r ON r.id = g.request_id
                WHERE r.requester_id = %s AND r.person_id = %s AND r.status = 'approved'
            """, (user_id, person_id))
            granted_ids = [r["note_id"] for r in cur.fetchall()]

            cur.execute("""
                SELECT pn.id, pn.note, pn.created_at, pn.visibility, u.username
                FROM person_notes pn
                JOIN users u ON u.id = pn.user_id
                WHERE pn.person_id = %s AND (pn.visibility = 'all' OR pn.id = ANY(%s))
                ORDER BY pn.created_at DESC
            """, (person_id, granted_ids or [0]))
            notes = [dict(r) for r in cur.fetchall()]

            cur.execute("""
                SELECT COUNT(*) as count FROM person_notes
                WHERE person_id = %s AND visibility = 'admin'
                AND id != ALL(%s)
            """, (person_id, granted_ids or [0]))
            hidden_count = cur.fetchone()["count"]

    return jsonify({"notes": notes, "hidden_count": hidden_count})

@app.route("/api/people/<int:person_id>/notes", methods=["POST"])
@login_required
def save_person_notes(person_id):
    data = request.json or {}
    note = data.get("notes", "").strip()
    visibility = data.get("visibility", "all")
    if not note:
        return jsonify({"error": "no note provided"}), 400
    if session.get("role") not in ("admin", "superadmin"):
        visibility = "all"
    with db_ctx() as (conn, cur):
        cur.execute("""
            INSERT INTO person_notes (person_id, user_id, note, created_at, visibility)
            VALUES (%s, %s, %s, %s, %s)
        """, (person_id, session["user_id"], note, datetime.now().isoformat(), visibility))
        conn.commit()
    return jsonify({"ok": True})

@app.route("/api/people/search")
@login_required
def api_people_search():
    q = request.args.get("q", "").strip()
    field = request.args.get("field", "notes")
    if not q:
        return jsonify([])

    with db_ctx() as (_, cur):
        if field == "notes":
            if session.get("role") in ("admin", "superadmin"):
                cur.execute("""
                    SELECT DISTINCT p.id, p.name, p.role, p.phones, p.emails,
                        c.name as company_name,
                        pn.note as matched_note, pn.created_at, u.username
                    FROM people p
                    JOIN companies c ON c.id = p.company_id
                    JOIN person_notes pn ON pn.person_id = p.id
                    JOIN users u ON u.id = pn.user_id
                    WHERE pn.note ILIKE %s
                    ORDER BY p.name
                """, (f"%{q}%",))
            else:
                cur.execute("""
                    SELECT DISTINCT p.id, p.name, p.role, p.phones, p.emails,
                        c.name as company_name,
                        pn.note as matched_note, pn.created_at, u.username
                    FROM people p
                    JOIN companies c ON c.id = p.company_id
                    JOIN person_notes pn ON pn.person_id = p.id
                    JOIN users u ON u.id = pn.user_id
                    WHERE pn.note ILIKE %s AND pn.visibility = 'all'
                    ORDER BY p.name
                """, (f"%{q}%",))
        elif field == "people":
            cur.execute("""
                SELECT p.id, p.name, p.role, p.phones, p.emails,
                       c.name as company_name,
                       NULL as matched_note, NULL as created_at, NULL as username
                FROM people p
                JOIN companies c ON c.id = p.company_id
                WHERE p.name ILIKE %s
                ORDER BY p.name
            """, (f"%{q}%",))
        elif field == "products":
            cur.execute("""
                SELECT DISTINCT p.id, p.name, p.role, p.phones, p.emails,
                       c.name as company_name,
                       e.products as matched_note, NULL as created_at, NULL as username
                FROM people p
                JOIN companies c ON c.id = p.company_id
                JOIN enrichment e ON e.contact_id = c.id
                WHERE e.products ILIKE %s
                ORDER BY p.name
            """, (f"%{q}%",))
        elif field == "company_notes":
            is_admin = session.get("role") in ("admin", "superadmin")
            if is_admin:
                cur.execute("""
                    SELECT DISTINCT ON (c.id)
                        p.id, p.name, p.role, p.phones, p.emails,
                        c.id as company_id, c.name as company_name,
                        cn.note as matched_note, cn.created_at, u.username
                    FROM people p
                    JOIN companies c ON c.id = p.company_id
                    JOIN company_notes cn ON cn.company_id = c.id
                    JOIN users u ON u.id = cn.user_id
                    WHERE cn.note ILIKE %s
                    ORDER BY c.id, p.name
                """, (f"%{q}%",))
            else:
                cur.execute("""
                    SELECT DISTINCT ON (c.id)
                        p.id, p.name, p.role, p.phones, p.emails,
                        c.id as company_id, c.name as company_name,
                        cn.note as matched_note, cn.created_at, u.username
                    FROM people p
                    JOIN companies c ON c.id = p.company_id
                    JOIN company_notes cn ON cn.company_id = c.id
                    JOIN users u ON u.id = cn.user_id
                    WHERE cn.note ILIKE %s AND cn.visibility = 'all'
                    ORDER BY c.id, p.name
                """, (f"%{q}%",))

        rows = cur.fetchall()

    return jsonify([{
        "id": r["id"],
        "company_id": r.get("company_id"),
        "name": r["name"],
        "role": r["role"],
        "phones": safe_json(r["phones"]),
        "emails": safe_json(r["emails"]),
        "company_name": r["company_name"],
        "matched_note": r["matched_note"],
        "created_at": r["created_at"],
        "username": r["username"]
    } for r in rows])


@app.route("/api/notes/<int:note_id>", methods=["PATCH"])
@login_required
def edit_note(note_id):
    data = request.json or {}
    note = data.get("note", "").strip()
    visibility = data.get("visibility")

    with db_ctx() as (conn, cur):
        cur.execute("SELECT user_id, visibility FROM person_notes WHERE id = %s", (note_id,))
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "note not found"}), 404
        if row["user_id"] != session["user_id"]:
            return jsonify({"error": "not your note"}), 403
        if not note:
            cur.execute("DELETE FROM note_access_grants WHERE note_id = %s", (note_id,))
            cur.execute("DELETE FROM person_notes WHERE id = %s", (note_id,))
        else:
            new_visibility = visibility if visibility and session.get("role") in ("admin", "superadmin") else row["visibility"]
            cur.execute("UPDATE person_notes SET note = %s, visibility = %s WHERE id = %s",
                        (note, new_visibility, note_id))
        conn.commit()
    return jsonify({"ok": True, "deleted": not note})

@app.route("/api/notes/cleanup", methods=["POST"])
@login_required
def cleanup_note():
    note = (request.json or {}).get("note", "").strip()
    if not note:
        return jsonify({"error": "no note provided"}), 400
    import openai
    client = openai.OpenAI()
    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{
            "role": "user",
            "content": f"""You are cleaning up a voice note from a business meeting.
            Summarize and condense the following note, fix grammar and punctuation, and make it concise and professional.
            Keep all important information but remove filler words and repetition.
            Return only the cleaned note, no explanation.

            Raw note:
            {note}"""
                    }]
                )
    return jsonify({"cleaned": response.choices[0].message.content.strip()})

# ─────────────────────────────────────────────────────────────────────────────
#  Pipeline execution + SSE streaming
# ─────────────────────────────────────────────────────────────────────────────

def _stream_process(job_id: str, cmd: list[str], cwd: str = None):
    """Run a subprocess and buffer its output for SSE."""
    with _process_lock:
        _processes[job_id] = {"output": [], "done": False, "proc": None}

    try:
        proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            cwd=cwd or str(Path(__file__).parent),
            bufsize=1,
        )
        with _process_lock:
            _processes[job_id]["proc"] = proc

        for line in proc.stdout:
            with _process_lock:
                _processes[job_id]["output"].append(line.rstrip())

        proc.wait()
    except Exception as e:
        with _process_lock:
            _processes[job_id]["output"].append(f"ERROR: {e}")
    finally:
        with _process_lock:
            _processes[job_id]["done"] = True


@app.route("/api/pipeline/start", methods=["POST"])
def start_pipeline():
    data = request.json or {}
    pipeline = data.get("pipeline")  # "ocr" | "enrich" | "normalize"
    job_id = str(uuid.uuid4())[:8]

    if pipeline == "ocr":
        folder = data.get("folder", "")
        if not folder:
            return jsonify({"error": "folder required"}), 400
        cmd = [sys.executable, "script.py", folder]
    elif pipeline == "enrich":
        cmd = [sys.executable, "enrich2.py"]
    elif pipeline == "normalize":
        cmd = [sys.executable, "normalizer.py"]
    else:
        return jsonify({"error": "unknown pipeline"}), 400

    t = threading.Thread(target=_stream_process, args=(job_id, cmd), daemon=True)
    t.start()
    return jsonify({"job_id": job_id})


@app.route("/api/pipeline/output/<job_id>")
def pipeline_output(job_id):
    """SSE endpoint — streams buffered + new output lines."""

    def generate():
        sent = 0
        while True:
            with _process_lock:
                job = _processes.get(job_id)
            if not job:
                yield "data: Job not found\n\n"
                return

            with _process_lock:
                lines = job["output"][sent:]
                done = job["done"]

            for line in lines:
                yield f"data: {line}\n\n"
            sent += len(lines)

            if done and sent >= len(_processes.get(job_id, {}).get("output", [])):
                yield "data: __DONE__\n\n"
                return

            time.sleep(0.3)

    return Response(
        stream_with_context(generate()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )

# ─────────────────────────────────────────────────────────────────────────────
#  Access requests
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/api/people/<int:person_id>/access-request", methods=["POST"])
@login_required
def request_access(person_id):
    user_id = session["user_id"]
    with db_ctx() as (conn, cur):
        cur.execute("""
            SELECT id FROM note_access_requests
            WHERE requester_id = %s AND person_id = %s AND status = 'pending'
        """, (user_id, person_id))
        if cur.fetchone():
            return jsonify({"error": "already requested"}), 409
        cur.execute("""
            INSERT INTO note_access_requests (requester_id, person_id, status, created_at)
            VALUES (%s, %s, 'pending', %s)
        """, (user_id, person_id, datetime.now().isoformat()))
        conn.commit()
    return jsonify({"ok": True})


@app.route("/api/people/<int:person_id>/access-request", methods=["GET"])
@login_required
def get_access_request_status(person_id):
    user_id = session["user_id"]
    with db_ctx() as (_, cur):
        cur.execute("""
            SELECT r.id, r.status,
                ARRAY_AGG(g.note_id) FILTER (WHERE g.note_id IS NOT NULL) as granted_note_ids
            FROM note_access_requests r
            LEFT JOIN note_access_grants g ON g.request_id = r.id
            WHERE r.requester_id = %s AND r.person_id = %s
            GROUP BY r.id, r.status
            ORDER BY r.created_at DESC
            LIMIT 1
        """, (user_id, person_id))
        row = cur.fetchone()

    if not row or not row["id"]:
        return jsonify({"status": "none"})
    return jsonify({
        "status": row["status"],
        "request_id": row["id"],
        "granted_note_ids": row["granted_note_ids"] or []
    })


@app.route("/api/admin/access-requests", methods=["GET"])
@admin_required
def get_access_requests():
    with db_ctx() as (_, cur):
        cur.execute("""
            SELECT r.id, r.status, r.created_at, r.person_id,
                   u.username as requester,
                   p.name as person_name,
                   c.name as company_name
            FROM note_access_requests r
            JOIN users u ON u.id = r.requester_id
            JOIN people p ON p.id = r.person_id
            JOIN companies c ON c.id = p.company_id
            WHERE r.status = 'pending'
            ORDER BY r.created_at DESC
        """)
        rows = cur.fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/admin/access-requests/<int:request_id>/notes", methods=["GET"])
@admin_required
def get_request_notes(request_id):
    with db_ctx() as (_, cur):
        cur.execute("SELECT person_id, requester_id FROM note_access_requests WHERE id = %s", (request_id,))
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "not found"}), 404

        cur.execute("""
            SELECT g.note_id FROM note_access_grants g
            JOIN note_access_requests r ON r.id = g.request_id
            WHERE r.requester_id = %s AND r.person_id = %s AND r.status = 'approved'
        """, (row["requester_id"], row["person_id"]))
        already_granted = [r["note_id"] for r in cur.fetchall()]

        cur.execute("""
            SELECT pn.id, pn.note, pn.created_at, pn.visibility, u.username
            FROM person_notes pn
            JOIN users u ON u.id = pn.user_id
            WHERE pn.person_id = %s AND pn.visibility = 'admin'
            AND pn.id != ALL(%s)
            ORDER BY pn.created_at DESC
        """, (row["person_id"], already_granted or [0]))
        notes = cur.fetchall()
    return jsonify([dict(n) for n in notes])


@app.route("/api/admin/access-requests/<int:request_id>/resolve", methods=["POST"])
@admin_required
def resolve_access_request(request_id):
    data = request.json or {}
    action = data.get("action")
    note_ids = data.get("note_ids", [])

    if action not in ("approve", "deny"):
        return jsonify({"error": "invalid action"}), 400

    status = "approved" if action == "approve" else "denied"
    with db_ctx() as (conn, cur):
        cur.execute("""
            UPDATE note_access_requests SET status = %s, resolved_at = %s WHERE id = %s
        """, (status, datetime.now().isoformat(), request_id))
        if action == "approve" and note_ids:
            for note_id in note_ids:
                cur.execute(
                    "INSERT INTO note_access_grants (request_id, note_id) VALUES (%s, %s)",
                    (request_id, note_id)
                )
        conn.commit()
    return jsonify({"ok": True})


# ─────────────────────────────────────────────────────────────────────────────
#  Company visibility (admin)
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/api/admin/companies/visibility", methods=["GET"])
@admin_required
def api_admin_get_visibility():
    with db_ctx() as (_, cur):
        cur.execute("""
            SELECT c.id, c.name,
                   COALESCE(v.hide_company, FALSE) AS hide_company,
                   COALESCE(v.hide_employees, FALSE) AS hide_employees,
                   COALESCE(v.hide_contact_info, FALSE) AS hide_contact_info
            FROM companies c
            LEFT JOIN company_visibility v ON v.company_id = c.id
            ORDER BY c.name
        """)
        rows = cur.fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/company/<int:company_id>/visibility", methods=["POST"])
@admin_required
def api_set_visibility(company_id):
    data = request.json or {}
    hide_company = bool(data.get("hide_company", False))
    hide_employees = bool(data.get("hide_employees", False))
    hide_contact_info = bool(data.get("hide_contact_info", False))

    with db_ctx() as (conn, cur):
        cur.execute("""
            INSERT INTO company_visibility (company_id, hide_company, hide_employees, hide_contact_info)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (company_id) DO UPDATE
                SET hide_company = EXCLUDED.hide_company,
                    hide_employees = EXCLUDED.hide_employees,
                    hide_contact_info = EXCLUDED.hide_contact_info
        """, (company_id, hide_company, hide_employees, hide_contact_info))
        conn.commit()
    return jsonify({"ok": True})

# ─────────────────────────────────────────────────────────────────────────────
#  Admin record editing
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/api/admin/company/<int:company_id>", methods=["GET"])
@admin_required
def api_admin_get_company(company_id):
    with db_ctx() as (_, cur):
        cur.execute("SELECT * FROM companies WHERE id = %s", (company_id,))
        company = cur.fetchone()
        if not company:
            return jsonify({"error": "not found"}), 404
        cur.execute("SELECT * FROM enrichment WHERE contact_id = %s", (company_id,))
        enrichment = cur.fetchone()
        cur.execute("SELECT category FROM company_categories WHERE company_id = %s ORDER BY category", (company_id,))
        categories = [r["category"] for r in cur.fetchall()]

    return jsonify({
        "id":               company["id"],
        "name":             company["name"] or "",
        "phones":           safe_json(company["phones"]),
        "emails":           safe_json(company.get("emails")),
        "addresses":        safe_json(company["addresses"]),
        "websites":         safe_json(company["websites"]),
        "primary_industry": enrichment["primary_industry"] if enrichment else "",
        "sub_industry":     enrichment["sub_industry"] if enrichment else "",
        "company_type":     safe_json(enrichment["company_type"]) if enrichment else [],
        "products":         safe_json(enrichment["products"]) if enrichment else [],
        "markets":          safe_json(enrichment["markets"]) if enrichment else [],
        "categories":       categories,
    })


@app.route("/api/admin/company/<int:company_id>", methods=["PATCH"])
@admin_required
def api_admin_update_company(company_id):
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "name cannot be empty"}), 400

    with db_ctx() as (conn, cur):
        cur.execute("""
            UPDATE companies
            SET name = %s, phones = %s, emails = %s, addresses = %s, websites = %s
            WHERE id = %s
        """, (
            name,
            json.dumps(data.get("phones", [])),
            json.dumps(data.get("emails", [])),
            json.dumps(data.get("addresses", [])),
            json.dumps(data.get("websites", [])),
            company_id,
        ))

        cur.execute("SELECT id FROM enrichment WHERE contact_id = %s", (company_id,))
        if cur.fetchone():
            cur.execute("""
                UPDATE enrichment
                SET primary_industry = %s, sub_industry = %s,
                    company_type = %s, products = %s, markets = %s
                WHERE contact_id = %s
            """, (
                data.get("primary_industry") or None,
                data.get("sub_industry") or None,
                json.dumps(data.get("company_type", [])),
                json.dumps(data.get("products", [])),
                json.dumps(data.get("markets", [])),
                company_id,
            ))
        else:
            cur.execute("""
                INSERT INTO enrichment (contact_id, primary_industry, sub_industry, company_type, products, markets)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (
                company_id,
                data.get("primary_industry") or None,
                data.get("sub_industry") or None,
                json.dumps(data.get("company_type", [])),
                json.dumps(data.get("products", [])),
                json.dumps(data.get("markets", [])),
            ))
        conn.commit()

    save_categories(company_id, data.get("categories", []))
    return jsonify({"ok": True})


@app.route("/api/admin/person/<int:person_id>", methods=["PATCH"])
@admin_required
def api_admin_update_person(person_id):
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "name cannot be empty"}), 400

    with db_ctx() as (conn, cur):
        cur.execute("""
            UPDATE people
            SET name = %s, role = %s, phones = %s, emails = %s
            WHERE id = %s
        """, (
            name,
            data.get("role") or None,
            json.dumps(data.get("phones", [])),
            json.dumps(data.get("emails", [])),
            person_id,
        ))
        conn.commit()
    return jsonify({"ok": True})

# ─────────────────────────────────────────────────────────────────────────────
#  Company notes
# ─────────────────────────────────────────────────────────────────────────────

@app.route("/api/company/<int:company_id>/notes", methods=["GET"])
@login_required
def get_company_notes(company_id):
    role = session.get("role")
    with db_ctx() as (_, cur):
        if role in ("admin", "superadmin"):
            cur.execute("""
                SELECT cn.id, cn.note, cn.note_type, cn.created_at, cn.visibility, u.username
                FROM company_notes cn
                JOIN users u ON u.id = cn.user_id
                WHERE cn.company_id = %s
                ORDER BY cn.created_at DESC
            """, (company_id,))
            notes = [dict(r) for r in cur.fetchall()]
            hidden_count = 0
        else:
            cur.execute("""
                SELECT cn.id, cn.note, cn.note_type, cn.created_at, cn.visibility, u.username
                FROM company_notes cn
                JOIN users u ON u.id = cn.user_id
                WHERE cn.company_id = %s AND cn.visibility = 'all'
                ORDER BY cn.created_at DESC
            """, (company_id,))
            notes = [dict(r) for r in cur.fetchall()]
            cur.execute("""
                SELECT COUNT(*) as count FROM company_notes
                WHERE company_id = %s AND visibility = 'admin'
            """, (company_id,))
            hidden_count = cur.fetchone()["count"]
    return jsonify({"notes": notes, "hidden_count": hidden_count})


@app.route("/api/company/<int:company_id>/notes", methods=["POST"])
@login_required
def save_company_note(company_id):
    data = request.json or {}
    note = data.get("note", "").strip()
    note_type = data.get("note_type", "general")
    visibility = data.get("visibility", "all")
    if not note:
        return jsonify({"error": "no note provided"}), 400
    if note_type not in ("intel", "general"):
        note_type = "general"
    if session.get("role") not in ("admin", "superadmin"):
        visibility = "all"
    with db_ctx() as (conn, cur):
        cur.execute("""
            INSERT INTO company_notes (company_id, user_id, note, note_type, visibility, created_at)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (company_id, session["user_id"], note, note_type, visibility, datetime.now().isoformat()))
        conn.commit()
    return jsonify({"ok": True})


@app.route("/api/notes/company/<int:note_id>", methods=["PATCH"])
@login_required
def edit_company_note(note_id):
    data = request.json or {}
    note = data.get("note", "").strip()
    with db_ctx() as (conn, cur):
        cur.execute("SELECT user_id FROM company_notes WHERE id = %s", (note_id,))
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "note not found"}), 404
        if row["user_id"] != session["user_id"]:
            return jsonify({"error": "not your note"}), 403
        if not note:
            cur.execute("DELETE FROM company_notes WHERE id = %s", (note_id,))
        else:
            cur.execute("UPDATE company_notes SET note = %s WHERE id = %s", (note, note_id))
        conn.commit()
    return jsonify({"ok": True, "deleted": not note})


if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port=5050)